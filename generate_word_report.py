import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_table(table, header_bg="1E3A8A", alt_bg="F8FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
    
    # Body rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if r_idx % 2 == 0 else "FFFFFF"
        for cell in row.cells:
            if bg != "FFFFFF":
                set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(30, 41, 59)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # #1E3A8A
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(14, 116, 144) # #0E7490
    return p

def add_callout(doc, text_list, title="💡 統計洞察與結論：", color_hex="EFF6FF", border_hex="2563EB"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, color_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(title + "\n")
    r_title.bold = True
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(29, 78, 216)
    
    for item in text_list:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(2)
        r_item = p_item.add_run("• " + item)
        r_item.font.size = Pt(9.5)
        r_item.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph()

doc = Document()

# Set standard page margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Set Default Font
style = doc.styles['Normal']
font = style.font
font.name = 'PingFang TC'
font.size = Pt(10.5)
font.color.rgb = RGBColor(51, 65, 85)

# --- Title Section ---
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(4)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
t_run = title_p.add_run("123 資料夾統計數據深度分析報告")
t_run.font.size = Pt(24)
t_run.font.bold = True
t_run.font.color.rgb = RGBColor(30, 58, 138)

sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_before = Pt(0)
sub_p.paragraph_format.space_after = Pt(16)
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
s_run = sub_p.add_run("涵蓋 Kruskal-Wallis 檢定、Mann-Whitney U 檢定、事後成對比較與專業圖表繪製")
s_run.font.size = Pt(12)
s_run.font.color.rgb = RGBColor(100, 116, 139)

meta_p = doc.add_paragraph()
meta_p.paragraph_format.space_after = Pt(16)
meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
m_run = meta_p.add_run("分析目錄：/Users/chenyichieh/Downloads/123   |   分析檔案數：3 份 Excel   |   報告日期：2026 年 10 月")
m_run.font.size = Pt(9)
m_run.font.color.rgb = RGBColor(148, 163, 184)

# ========================================================
# 1. 執行摘要
# ========================================================
add_heading_1(doc, "一、執行摘要與綜合成果總覽 (Executive Summary Dashboard)")

p = doc.add_paragraph("本專案針對 123 資料夾內的三個統計學數據檔案進行全面探索性資料分析 (EDA)、母數與非母數假設檢定、效應量 (Effect Size) 評估與視覺化圖表生成。下圖為三項研究之中位數與檢定結論之綜合儀表板：")
p.paragraph_format.space_after = Pt(8)

if os.path.exists("figure4_summary_dashboard.png"):
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_before = Pt(6)
    pic_p.paragraph_format.space_after = Pt(4)
    run_img = pic_p.add_run()
    run_img.add_picture("figure4_summary_dashboard.png", width=Inches(6.2))
    
    cap_p = doc.add_paragraph("圖 1：三項統計研究綜合對比儀表板 (Summary Dashboard)")
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.runs[0].font.size = Pt(8.5)
    cap_p.runs[0].font.italic = True
    cap_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    cap_p.paragraph_format.space_after = Pt(12)

# Executive Table
headers1 = ["檔案名稱", "研究主題", "樣本數", "推薦檢定方法", "統計檢定量", "p-value", "顯著性", "核心結論"]
data1 = [
    ["KW_Blood type.xlsx", "血型與身高相關性", "N = 44", "Kruskal-Wallis\n(ANOVA)", "H = 1.507\n(F = 0.460)", "0.6806\n(0.7118)", "不顯著 (ns)", "血型與身高無顯著關聯，符合現代遺傳學常理。"],
    ["KW_Pro_salary.xlsx", "四大職業運動年薪", "N = 40", "Kruskal-Wallis\n(Dunn's post-hoc)", "H = 27.611", "4.38 × 10⁻⁶", "極顯著 (***)", "WNBA 與男子三大聯盟存在 234 倍性別薪資落差。"],
    ["Mann whitney.xlsx", "乳癌 vs 肺癌 HER2", "N = 17", "Mann-Whitney U", "U = 71.0\n(W = 107.0)", "8.88 × 10⁻⁴", "極顯著 (***)", "乳癌 HER2 顯著高於肺腺癌 (Cliff's δ = 0.972)。"]
]

t1 = doc.add_table(rows=len(data1)+1, cols=len(headers1))
for j, h in enumerate(headers1):
    t1.rows[0].cells[j].paragraphs[0].text = h
for i, row in enumerate(data1):
    for j, val in enumerate(row):
        t1.rows[i+1].cells[j].paragraphs[0].text = val
style_table(t1)
doc.add_paragraph()

# ========================================================
# 2. 研究一
# ========================================================
add_heading_1(doc, "二、研究一：KW_Blood type.xlsx — 血型與身高關聯性分析")

p = doc.add_paragraph("本檔案探索 ABO 血型分類 (A, B, O, AB) 是否對受試者之成人身高 (Height, cm) 造成統計學上的系統性差異。自變數為 4 個血型組別，應變數為連續數值型身高。")

add_heading_2(doc, "1. 敘述統計與常態性檢定")
headers_bt = ["血型 (Blood Type)", "樣本數 (n)", "平均數 ± 標準差 (cm)", "中位數 (cm)", "四分位距 IQR", "範圍 [Min - Max]", "Shapiro-Wilk 檢定"]
data_bt = [
    ["A 型", "12", "174.00 ± 7.63", "173.50", "7.50", "[158, 185]", "W = 0.941, p = 0.4807 (常態)"],
    ["B 型", "10", "171.50 ± 9.07", "173.00", "9.50", "[157, 184]", "W = 0.948, p = 0.5969 (常態)"],
    ["O 型", "9", "170.44 ± 11.31", "170.00", "11.00", "[154, 187]", "W = 0.942, p = 0.6061 (常態)"],
    ["AB 型", "13", "169.92 ± 9.11", "172.00", "9.00", "[154, 185]", "W = 0.984, p = 0.9696 (常態)"],
    ["整體總計", "44", "171.50 ± 9.05", "172.00", "11.00", "[154, 187]", "W = 0.983, p = 0.7416 (常態)"]
]
t_bt = doc.add_table(rows=len(data_bt)+1, cols=len(headers_bt))
for j, h in enumerate(headers_bt):
    t_bt.rows[0].cells[j].paragraphs[0].text = h
for i, row in enumerate(data_bt):
    for j, val in enumerate(row):
        t_bt.rows[i+1].cells[j].paragraphs[0].text = val
style_table(t_bt)

add_heading_2(doc, "2. 假設檢定實證結果")
p = doc.add_paragraph()
p.add_run("• 變異數同質性檢定 (Levene's Test)：").bold = True
p.add_run("W = 0.4563, p = 0.7143 (p > 0.05，滿足變異數齊一性假設)。\n")
p.add_run("• Kruskal-Wallis 非母數檢定：").bold = True
p.add_run("H = 1.5070, df = 3, p = 0.6806 (p > 0.05，無顯著差異 ns)。\n")
p.add_run("• 單因子變異數分析 (One-way ANOVA)：").bold = True
p.add_run("F(3, 40) = 0.4600, p = 0.7118 (p > 0.05，無顯著差異 ns)。\n")
p.add_run("• 效應量評估：").bold = True
p.add_run("η² = 0.033, ε² = 0.037 (效應量微小，組間差異僅佔總變異的 3.3%)。")

if os.path.exists("figure1_blood_type_height.png"):
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_before = Pt(8)
    pic_p.paragraph_format.space_after = Pt(4)
    run_img = pic_p.add_run()
    run_img.add_picture("figure1_blood_type_height.png", width=Inches(6.0))
    
    cap_p = doc.add_paragraph("圖 2：血型與身高分佈箱型圖與機率密度估計圖 (KDE)")
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.runs[0].font.size = Pt(8.5)
    cap_p.runs[0].font.italic = True
    cap_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)

add_callout(doc, [
    "檢定結論：無母數 Kruskal-Wallis 與母數 ANOVA 檢定結果完全一致，p 值均遠大於 0.05，無法拒絕虛無假說 (Fail to reject H₀)。",
    "生理遺傳學意義：ABO 血型由第 9 號染色體之單一基因座決定，而人類身高級數屬於典型多基因遺傳 (Polygenic inheritance) 並受後天營養環境影響，兩者機轉獨立，無顯著差異符合科學事實。",
    "附註說明：該檔案工作表 2 包含 10 筆 NFL 球員年薪，推測為建檔時之暫存資料，主分析以工作表 1 為準。"
])

# ========================================================
# 3. 研究二
# ========================================================
add_heading_1(doc, "三、研究二：KW_Pro_salary.xlsx — 職業運動年薪結構與性別差距研究")

p = doc.add_paragraph("本檔案取樣四大職業運動聯盟各 10 位頂尖選手 (總樣本數 N = 40)：NBA (男子職籃)、MLB (男子職棒)、NFL (美式足球) 及 WNBA (女子職籃)，針對年薪、性別與年齡結構進行深度剖析。")

add_heading_2(doc, "1. 聯盟年薪敘述統計")
headers_sal = ["聯盟 / 運動項目", "性別", "樣本數", "平均年薪 (Mean ± SD)", "中位數年薪 (Median)", "最高年薪代表 (Max)", "最低年薪代表 (Min)"]
data_sal = [
    ["NBA (男籃)", "男 (1)", "10", "$45,464,546 ± $7,037,160", "$46,623,717", "Stephen Curry ($55.76M)", "Jayson Tatum ($34.85M)"],
    ["MLB (棒球)", "男 (1)", "10", "$42,370,833 ± $11,121,253", "$38,500,000", "Shohei Ohtani 大谷翔平 ($70.00M)", "Corey Seager ($32.50M)"],
    ["NFL (美足)", "男 (1)", "10", "$52,460,000 ± $4,408,880", "$52,750,000", "Dak Prescott ($60.00M)", "Patrick Mahomes ($45.00M)"],
    ["WNBA (女籃)", "女 (2)", "10", "$199,922 ± $47,436", "$203,500", "Jackie Young ($252,450)", "Caitlin Clark ($76,535)"]
]
t_sal = doc.add_table(rows=len(data_sal)+1, cols=len(headers_sal))
for j, h in enumerate(headers_sal):
    t_sal.rows[0].cells[j].paragraphs[0].text = h
for i, row in enumerate(data_sal):
    for j, val in enumerate(row):
        t_sal.rows[i+1].cells[j].paragraphs[0].text = val
style_table(t_sal)

add_heading_2(doc, "2. 假說檢定與事後成對比較 (Dunn's Post-Hoc Test)")
p = doc.add_paragraph()
p.add_run("• 分布特徵與適格性：").bold = True
p.add_run("MLB (p = 0.0083) 與 WNBA (p = 0.0009) 具顯著偏態，且變異數極不均齊 (Levene p = 0.020)，故必須採用非參數 Kruskal-Wallis 檢定。\n")
p.add_run("• Kruskal-Wallis 主檢定：").bold = True
p.add_run("H = 27.6106, df = 3, p = 4.384 × 10⁻⁶ (p < 0.0001，達極高度顯著差異 ***)。\n")
p.add_run("• 效應量：").bold = True
p.add_run("ε² = 0.684 (極大效應量，聯盟分類解釋了絕大部分薪資變異)。")

headers_dunn = ["成對比較組別", "平均秩差 (Rank Diff)", "標準化檢定量 Z", "未校正 p-value", "Bonferroni 校正後 p 值", "顯著性判定"]
data_dunn = [
    ["NFL vs WNBA", "+26.75", "5.117", "3.11 × 10⁻⁷", "1.87 × 10⁻⁶", "極顯著 (***)"],
    ["NBA vs WNBA", "+18.70", "3.577", "3.48 × 10⁻⁴", "2.09 × 10⁻³", "高度顯著 (**)"],
    ["MLB vs WNBA", "+14.55", "2.783", "5.39 × 10⁻³", "3.23 × 10⁻²", "顯著 (*)"],
    ["NFL vs MLB", "+12.20", "2.334", "0.0196", "0.1177", "不顯著 (趨勢)"],
    ["NFL vs NBA", "+8.05", "1.540", "0.1236", "0.7417", "不顯著 (ns)"],
    ["NBA vs MLB", "+4.15", "0.794", "0.4273", "1.0000", "不顯著 (ns)"]
]
t_dunn = doc.add_table(rows=len(data_dunn)+1, cols=len(headers_dunn))
for j, h in enumerate(headers_dunn):
    t_dunn.rows[0].cells[j].paragraphs[0].text = h
for i, row in enumerate(data_dunn):
    for j, val in enumerate(row):
        t_dunn.rows[i+1].cells[j].paragraphs[0].text = val
style_table(t_dunn)

add_heading_2(doc, "3. 性別薪資落差與年齡趨勢剖析")
p = doc.add_paragraph()
p.add_run("• 性別運動薪資落差：").bold = True
p.add_run("男性選手 (n=30) 平均年薪為 4,677 萬美元，女性選手 (n=10) 平均年薪僅 20 萬美元，差距高達 233.9 倍！Mann-Whitney U 檢定 U = 300.0, p = 1.84 × 10⁻⁷ (***)。\n")
p.add_run("• 年齡與薪資關係：").bold = True
p.add_run("NBA 內部年齡與年薪呈顯著正相關 (Spearman rs = 0.675, p = 0.0322)，反映勞資協議中年資與超級頂薪 (Supermax) 條款之作用。")

if os.path.exists("figure2_pro_salary_analysis.png"):
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_before = Pt(8)
    pic_p.paragraph_format.space_after = Pt(4)
    run_img = pic_p.add_run()
    run_img.add_picture("figure2_pro_salary_analysis.png", width=Inches(6.2))
    
    cap_p = doc.add_paragraph("圖 3：四大職業運動年薪結構、性別差距與年齡趨勢圖")
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.runs[0].font.size = Pt(8.5)
    cap_p.runs[0].font.italic = True
    cap_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)

add_callout(doc, [
    "職業運動產業結構：男子三大聯盟 (NFL, NBA, MLB) 選手年薪主要落在 3,200 萬至 7,000 萬美元之間，組間無顯著統計差異；而 WNBA 頂尖選手最高僅約 25 萬美元，受限於聯盟轉播合約體量與硬工資帽制度。",
    "事後檢定證實：顯著性主要完全來自於 WNBA 與男子三大聯盟之間的龐大落差 (校正後 p < 0.05)。"
])

# ========================================================
# 4. 研究三
# ========================================================
add_heading_1(doc, "四、研究三：Mann whitney.xlsx — 乳腺癌 vs 肺腺癌之 HER2 表現量檢定")

p = doc.add_paragraph("本檔案研究臨床腫瘤組織切片中，乳腺癌 (Breast Cancer, BC, n=8) 與肺腺癌 (Lung Adenocarcinoma, LUAD, n=9) 之 HER2 (人表皮生長因子受體 2) 蛋白質表現量差異。")

add_heading_2(doc, "1. 敘述統計與分佈檢定")
headers_mw = ["癌症分類 (Annotation)", "樣本數 (n)", "平均數 ± 標準差", "中位數 (Median)", "四分位距 IQR", "範圍 [Min - Max]", "Shapiro-Wilk 常態性檢定"]
data_mw = [
    ["乳腺癌 (BC, Group 1)", "8", "434.50 ± 202.05", "446.50", "343.25", "[156, 666]", "W = 0.892, p = 0.2220 (常態)"],
    ["肺腺癌 (LUAD, Group 2)", "9", "128.44 ± 25.62", "118.00", "13.00", "[111, 187]", "W = 0.724, p = 0.0033 (顯著偏態)"]
]
t_mw = doc.add_table(rows=len(data_mw)+1, cols=len(headers_mw))
for j, h in enumerate(headers_mw):
    t_mw.rows[0].cells[j].paragraphs[0].text = h
for i, row in enumerate(data_mw):
    for j, val in enumerate(row):
        t_mw.rows[i+1].cells[j].paragraphs[0].text = val
style_table(t_mw)

add_heading_2(doc, "2. Mann-Whitney U 假設檢定與效應量評估")
p = doc.add_paragraph()
p.add_run("• 檢定適格性：").bold = True
p.add_run("LUAD 樣本具極顯著偏態 (p = 0.0033)，且 BC 的變異數為 LUAD 的 62 倍 (Levene p = 0.0051)，傳統獨立樣本 t 檢定假設完全被打破，必須使用無母數 Mann-Whitney U 檢定。\n")
p.add_run("• Mann-Whitney U 統計檢定：").bold = True
p.add_run("U = 71.0, W = 107.0, 雙尾 p = 8.8787 × 10⁻⁴ (p < 0.001 ***，達極高度顯著差異)。\n")
p.add_run("• 效應量 Cliff's Delta (δ)：").bold = True
p.add_run("δ = (71 - 1) / 72 = 0.9722 (接近 1.0 之極致大效應量)。\n")
p.add_run("• 常態語言效應量 (CLES)：").bold = True
p.add_run("隨機抽取一名乳癌病患，其 HER2 高於肺腺癌之機率高達 98.6%。")

if os.path.exists("figure3_mann_whitney_her2.png"):
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_before = Pt(8)
    pic_p.paragraph_format.space_after = Pt(4)
    run_img = pic_p.add_run()
    run_img.add_picture("figure3_mann_whitney_her2.png", width=Inches(6.0))
    
    cap_p = doc.add_paragraph("圖 4：乳腺癌與肺腺癌 HER2 表現量小提琴圖、箱型圖與 ECDF 累積分配函數")
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.runs[0].font.size = Pt(8.5)
    cap_p.runs[0].font.italic = True
    cap_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)

add_callout(doc, [
    "高度分離特徵：8 個 BC 樣本中僅有 1 筆數值 (156) 落入較低區間，其餘皆大於 250；而 9 個 LUAD 樣本中僅有 1 筆數值 (187) 略高，其餘均小於等於 153，兩組幾近完全分離。",
    "腫瘤分子生物學意義：HER2/ERBB2 是乳腺癌臨床一線關鍵致癌基因 (Oncogene)，約 15-20% 乳癌會發生基因擴增而高度表現，適用曲妥珠單抗 (Trastuzumab) 標靶治療；而肺腺癌主要致癌突變為 EGFR 與 KRAS，HER2 過度表現非常罕見，數據精準反映了此病理分型特點。"
])

# ========================================================
# 5. 方法論總結
# ========================================================
add_heading_1(doc, "五、統計分析方法論總結 (Statistical Methodology)")

headers_method = ["檢定方法", "適用資料型態", "對應母數統計方法", "本專案應用範例", "選用無母數之核心優勢"]
data_method = [
    ["Kruskal-Wallis 檢定", "≥ 3 組獨立連續或等級變數", "單因子變異數分析 (ANOVA)", "血型與身高、四大運動聯盟年薪", "對小樣本、偏態分配及極端離群值具極佳強韌性。"],
    ["Mann-Whitney U 檢定", "2 組獨立連續或等級變數", "獨立樣本雙尾 t 檢定 (t-test)", "乳腺癌 vs 肺腺癌 HER2 表現", "以秩和為基礎，不依賴常態分佈與等分散性假設。"],
    ["Dunn's 事後檢定", "KW 檢定顯著後之成對檢定", "Tukey HSD / Scheffé 檢定", "四大運動聯盟成對年薪差異", "以總秩次標準差進行 Z 檢定，搭配 Bonferroni 校正控制型一誤差。"]
]
t_m = doc.add_table(rows=len(data_method)+1, cols=len(headers_method))
for j, h in enumerate(headers_method):
    t_m.rows[0].cells[j].paragraphs[0].text = h
for i, row in enumerate(data_method):
    for j, val in enumerate(row):
        t_m.rows[i+1].cells[j].paragraphs[0].text = val
style_table(t_m)

output_filename = "123資料夾數據統計分析報告.docx"
doc.save(output_filename)
print(f"Successfully generated {output_filename}!")
